#!/usr/bin/env python3
"""Build the Chinese RP1 ICSMD 2026 reading draft from the official Word template."""

from __future__ import annotations

import argparse
import re
from copy import deepcopy
from pathlib import Path

import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_TEMPLATE = ROOT / ".tmp" / "icsmd_paper" / "Paper_Template_transitional.docx"
DEFAULT_SOURCE = ROOT / "papers" / "ICSMD_2026" / "RP1_ICSMD2026_中文初稿_v1.md"
DEFAULT_OUTPUT = ROOT / "papers" / "ICSMD_2026" / "RP1_ICSMD2026_中文初稿_v1.docx"
TMP = ROOT / ".tmp" / "icsmd_paper" / "generated"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=40, start=50, bottom=40, end=50) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "bottom"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "8")
        tag.set(qn("w:color"), "000000")
    for edge in ("left", "right", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "nil")
    inside_h = borders.find(qn("w:insideH"))
    if inside_h is None:
        inside_h = OxmlElement("w:insideH")
        borders.append(inside_h)
    inside_h.set(qn("w:val"), "single")
    inside_h.set(qn("w:sz"), "2")
    inside_h.set(qn("w:color"), "BFBFBF")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, size: float | None = None, bold=None, italic=None, east="宋体") -> None:
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), east)
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def configure_style(doc: Document, name: str, size: float, east="宋体", *, bold=None, italic=None) -> None:
    style = doc.styles[name]
    style.font.name = "Times New Roman"
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), east)
    style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic


def add_inline_markdown(paragraph, text: str, size: float | None = None) -> None:
    text = clean_inline_math(text)
    pattern = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos : match.start()])
            set_run_font(run, size)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size, bold=True)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size, italic=True)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, size)


def clean_inline_math(text: str) -> str:
    """Convert the small LaTeX subset used by the Chinese draft to readable inline text."""
    replacements = {
        r"\alpha": "α",
        r"\eta": "η",
        r"\pi": "π",
        r"\ge": "≥",
        r"\in": "∈",
        r"\rightarrow": "→",
        r"\leftrightarrow": "↔",
    }
    for source, target in replacements.items():
        text = text.replace(source, target)
    text = text.replace(r"\(", "").replace(r"\)", "")
    text = re.sub(r"\\mathcal\s*\{?([A-Za-z])\}?", r"\1", text)
    text = re.sub(r"\\mathrm\{([^{}]+)\}", r"\1", text)
    text = re.sub(r"\\text\{([^{}]+)\}", r"\1", text)
    text = text.replace("{", "").replace("}", "")
    return text


def make_method_figure(path: Path) -> None:
    plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei", "DejaVu Sans"]
    plt.rcParams["axes.unicode_minus"] = False
    fig, ax = plt.subplots(figsize=(4.0, 5.0), dpi=220)
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")
    labels = [
        ("冻结文档划分", "15份构建文档 · 1934页\n开发/保留测试不进入主图"),
        ("页面级LLM抽取", "正文与同表同行组\n保留页码、原文、URL与哈希"),
        ("自动证据硬门控", "Schema · Domain/Range · E1/E2\n关系方向与支持 · 非推断"),
        ("CEPU审计图", "Claim与EvidenceAssertion分层\n合格 / 隔离 / 拒绝"),
        ("中文与来源治理", "中文术语发布门控\n来源族封顶佐证 · CQ评价"),
        ("中文发布图", "208条证据断言 · 203个Claim\n281个规范实体"),
    ]
    colors = ["#E7F0FA", "#E7F0FA", "#FFF1D6", "#E9E2F4", "#E2F3EA", "#DDEFE9"]
    ys = [0.89, 0.73, 0.57, 0.41, 0.25, 0.09]
    for idx, ((title, detail), y, color) in enumerate(zip(labels, ys, colors)):
        box = FancyBboxPatch(
            (0.08, y - 0.055), 0.84, 0.11,
            boxstyle="round,pad=0.008,rounding_size=0.015",
            linewidth=1.0, edgecolor="#305070", facecolor=color,
        )
        ax.add_patch(box)
        ax.text(0.50, y + 0.022, title, ha="center", va="center", fontsize=10.5, weight="bold", color="#17324D")
        ax.text(0.50, y - 0.022, detail, ha="center", va="center", fontsize=7.8, color="#263746")
        if idx < len(labels) - 1:
            ax.annotate("", xy=(0.50, y - 0.116), xytext=(0.50, y - 0.085), arrowprops=dict(arrowstyle="-|>", lw=1.2, color="#456B88"))
    fig.savefig(path, bbox_inches="tight", pad_inches=0.05, facecolor="white")
    plt.close(fig)


def crop_experiment_panels(source: Path, left: Path, right: Path) -> None:
    img = Image.open(source)
    w, h = img.size
    img.crop((0, 0, int(w * 0.515), h)).save(left)
    img.crop((int(w * 0.57), 0, w, h)).save(right)


def add_picture(doc: Document, path: Path, caption: str, width=Inches(3.35), alt="") -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    shape = p.add_run().add_picture(str(path), width=width)
    try:
        doc_pr = shape._inline.docPr
        doc_pr.set("descr", alt or caption)
        doc_pr.set("title", caption)
    except Exception:
        pass
    cap = doc.add_paragraph(style="figure caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    cap.paragraph_format.space_before = Pt(0)
    cap.paragraph_format.space_after = Pt(3)
    add_inline_markdown(cap, caption, 8.0)


def add_equation(doc: Document, raw: str) -> None:
    if "mathcal M" in raw:
        text = "M: D_build → (G_audit, G_release)                                      (1)"
    elif "c_i=" in raw:
        text = "c_i = (h_i, r_i, t_i),    a_i = ⟨c_i, e_i, π_i, u_i⟩                         (2)"
    elif "q_i=" in raw:
        text = "q_i = α_i · w_E(l_i) · w_R(η_i)                                     (3)"
    elif "H_{zh}" in raw:
        text = "H_zh(a_i) = H_S(a_i) T_zh(h_i) T_zh(t_i)                               (7)"
    elif "T_{zh}" in raw:
        text = "T_zh(x) = I_label I_term I_type I_status I_protect                         (6)"
    elif "H_S" in raw:
        text = "H_S(a_i) = I_sch I_d/r I_dir I_gnd I_ent I_E1/E2 I_noninf I_build I_family I_q≥0.8       (4)"
    elif "S_{fam}" in raw:
        text = "S_fam^(B)(c) = (1/B) Σ[j=1…min(B,|F_c|)] s_(j)(c)                       (5)"
    else:
        text = raw.replace("\\", "").replace("\n", " ")
    p = doc.add_paragraph(style="equation")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    set_run_font(run, 8.0, east="Cambria Math")
    run.font.name = "Cambria Math"


def add_table(doc: Document, rows: list[list[str]], table_no: int) -> None:
    captions = {
        1: "双图与自动门控分层统计",
        2: "固定20页提示契约服从与门控产出（时延单位：s）",
        3: "固定20页API调用的token消耗",
        4: "CQ v1 中没有合法发布路径的任务",
    }
    cap = doc.add_paragraph(style="table head")
    cap.paragraph_format.space_before = Pt(3)
    cap.paragraph_format.space_after = Pt(2)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline_markdown(cap, captions.get(table_no, f"表 {table_no}"), 8.0)
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table)
    for r_idx, values in enumerate(rows):
        row = table.rows[r_idx]
        if r_idx == 0:
            set_repeat_table_header(row)
        for c_idx, value in enumerate(values):
            cell = row.cells[c_idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if r_idx == 0:
                set_cell_shading(cell, "E7EEF5")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (r_idx == 0 or c_idx > 0) else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 0.9
            add_inline_markdown(p, value, 7.5 if table_no == 2 else 8.0)
            for run in p.runs:
                if r_idx == 0:
                    run.bold = True
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(0)
    after.paragraph_format.line_spacing = 0.25


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    i = start
    while i < len(lines) and lines[i].lstrip().startswith("|"):
        cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-+:?", c.replace(" ", "")) for c in cells):
            rows.append(cells)
        i += 1
    return rows, i


def build(template: Path, source: Path, output: Path) -> None:
    TMP.mkdir(parents=True, exist_ok=True)
    output.parent.mkdir(parents=True, exist_ok=True)
    method_fig = TMP / "method_pipeline_zh.png"
    ablation_fig = TMP / "ablation_panel_zh.png"
    replication_fig = TMP / "replication_panel_zh.png"
    make_method_figure(method_fig)
    crop_experiment_panels(
        ROOT / "results" / "experiments" / "research_point_1" / "b0_ours_ablation_v1" / "figures" / "ablation_and_replication.png",
        ablation_fig,
        replication_fig,
    )

    doc = Document(str(template))
    one_col = deepcopy(doc.sections[0]._sectPr)
    two_col = deepcopy(doc.sections[3]._sectPr)
    type_node = one_col.find(qn("w:type"))
    if type_node is None:
        type_node = OxmlElement("w:type")
        one_col.insert(0, type_node)
    type_node.set(qn("w:val"), "continuous")

    body = doc._element.body
    for child in list(body):
        body.remove(child)
    body.append(two_col)

    configure_style(doc, "paper title", 16, "微软雅黑", bold=True)
    configure_style(doc, "Author", 9.5, "宋体")
    configure_style(doc, "Abstract", 9.0, "宋体", bold=False)
    configure_style(doc, "Keywords", 9.0, "宋体", bold=False, italic=False)
    configure_style(doc, "Body Text", 10.0, "宋体")
    configure_style(doc, "Heading 1", 10.0, "宋体", bold=False)
    configure_style(doc, "Heading 2", 10.0, "宋体", italic=True)
    configure_style(doc, "figure caption", 8.0, "宋体")
    configure_style(doc, "table head", 8.0, "宋体")
    configure_style(doc, "references", 8.0, "宋体")
    doc.styles["Body Text"].paragraph_format.space_after = Pt(2.4)
    doc.styles["Body Text"].paragraph_format.line_spacing = 1.0
    doc.styles["Body Text"].paragraph_format.first_line_indent = Inches(0.15)
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(4)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(2)
    doc.styles["Heading 1"].paragraph_format.keep_with_next = True
    doc.styles["Heading 2"].paragraph_format.space_before = Pt(3)
    doc.styles["Heading 2"].paragraph_format.space_after = Pt(1)
    doc.styles["Heading 2"].paragraph_format.keep_with_next = True
    doc.styles["references"].paragraph_format.left_indent = Inches(0.15)
    doc.styles["references"].paragraph_format.first_line_indent = Inches(-0.15)
    doc.styles["references"].paragraph_format.space_after = Pt(1.5)

    lines = source.read_text(encoding="utf-8").splitlines()
    title = lines[0][2:].strip()
    english = lines[2].strip("*")
    authors = lines[4]
    affiliation = lines[6]
    draft_note = lines[8].lstrip("> ")

    p = doc.add_paragraph(style="paper title")
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(title)
    set_run_font(run, 16, bold=True, east="微软雅黑")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(english)
    set_run_font(run, 9.5, italic=True, east="微软雅黑")
    p = doc.add_paragraph(style="Author")
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(1)
    add_inline_markdown(p, authors, 9.2)
    p = doc.add_paragraph(style="Author")
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    add_inline_markdown(p, affiliation, 8.0)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(draft_note)
    set_run_font(run, 7.2, italic=True, east="宋体")
    run.font.color.rgb = RGBColor(90, 90, 90)

    section_break = doc.add_paragraph()
    section_break.paragraph_format.space_after = Pt(0)
    section_break._p.get_or_add_pPr().append(one_col)

    i = 9
    table_no = 0
    in_math = False
    math_lines: list[str] = []
    in_refs = False
    abstract_pending = False
    while i < len(lines):
        raw = lines[i]
        line = raw.strip()
        if i <= 9 or not line:
            i += 1
            continue
        if line == r"\[":
            in_math = True
            math_lines = []
            i += 1
            continue
        if in_math:
            if line == r"\]":
                add_equation(doc, " ".join(math_lines))
                in_math = False
            else:
                math_lines.append(line)
            i += 1
            continue
        if line.startswith("## "):
            heading = line[3:]
            numbered_heading = heading
            if heading == "摘要":
                abstract_pending = True
                i += 1
                continue
            if heading == "参考文献":
                p = doc.add_paragraph(style="Heading 5")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.page_break_before = False
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(2)
                add_inline_markdown(p, "参考文献", 8.8)
                in_refs = True
                i += 1
                continue
            heading = re.sub(r"^[IVX]+\.\s*", "", heading)
            p = doc.add_paragraph(style="Heading 1")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline_markdown(p, heading, 10.0)
            if numbered_heading.startswith("III."):
                add_picture(doc, method_fig, "CEPU证据图谱构建与双门控发布流程。", alt="从冻结文档划分到中文发布图的六阶段流程")
            i += 1
            continue
        if line.startswith("### "):
            heading = re.sub(r"^[A-Z]\.\s*", "", line[4:])
            p = doc.add_paragraph(style="Heading 2")
            add_inline_markdown(p, heading, 10.0)
            i += 1
            continue
        if line.startswith("|"):
            rows, i = parse_table(lines, i)
            table_no += 1
            add_table(doc, rows, table_no)
            continue
        if abstract_pending:
            p = doc.add_paragraph(style="Abstract")
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.first_line_indent = Inches(0.14)
            run = p.add_run("摘要—")
            set_run_font(run, 9.0, bold=True, italic=True)
            add_inline_markdown(p, line, 9.0)
            abstract_pending = False
            i += 1
            continue
        if line.startswith("**关键词—**"):
            p = doc.add_paragraph(style="Keywords")
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.first_line_indent = Inches(0.14)
            run = p.add_run("关键词—")
            set_run_font(run, 9.0, bold=True, italic=True)
            add_inline_markdown(p, line[len("**关键词—**") :].strip(), 9.0)
            i += 1
            continue
        if in_refs and re.match(r"^\[\d+\]", line):
            p = doc.add_paragraph(style="references")
            p.paragraph_format.keep_together = True
            line = re.sub(r"^\[\d+\]\s*", "", line)
            add_inline_markdown(p, line, 8.0)
            i += 1
            continue
        if re.match(r"^\d+\.\s", line):
            p = doc.add_paragraph(style="Body Text")
            p.paragraph_format.left_indent = Inches(0.16)
            p.paragraph_format.first_line_indent = Inches(-0.16)
            p.paragraph_format.space_after = Pt(1)
            add_inline_markdown(p, line, 10.0)
            i += 1
            continue
        if line.startswith(">"):
            i += 1
            continue
        p = doc.add_paragraph(style="Body Text")
        add_inline_markdown(p, line, 10.0)
        if line.startswith("在固定 8003 条审计候选上"):
            add_picture(doc, ablation_fig, "各门控在固定候选上的边际筛除作用。", alt="关系Schema、证据对齐、关系蕴含、溯源、分数和中文门控的边际筛除率")
        if line.startswith("真实语料中有 3 个精确 Claim"):
            add_picture(doc, replication_fig, "同来源族复制下的文档计数膨胀与来源族封顶不变性。", alt="同族复制后文档计数翻倍但来源族数与封顶指数不变")
        i += 1

    props = doc.core_properties
    props.title = title
    props.subject = "ICSMD 2026 Chinese reading draft for research point 1"
    props.author = "Authors to be confirmed"
    props.keywords = "marine pump; fault diagnosis; knowledge graph; provenance; Silver evidence"
    props.comments = "Chinese draft generated from the official ICSMD/IEEE Word template; author review required."

    # Remove all sample template footer/header content, including VML text boxes
    # that contain the placeholder IEEE copyright line.
    for section in doc.sections:
        section.different_first_page_header_footer = False
        for part in (section.header, section.footer):
            element = part._element
            for child in list(element):
                element.remove(child)
            element.append(OxmlElement("w:p"))

    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")
    doc.save(str(output))


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    build(args.template.resolve(), args.source.resolve(), args.output.resolve())
    print(args.output.resolve())


if __name__ == "__main__":
    main()
