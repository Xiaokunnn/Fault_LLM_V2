from __future__ import annotations

import re
import shutil
import sys
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
from matplotlib.patches import FancyBboxPatch
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[3]
PAPER_DIR = ROOT / "papers" / "ICSMD_2026"
SOURCE_MD = PAPER_DIR / "RP1_ICSMD2026_中文正式稿_v4.md"
TEMPLATE = PAPER_DIR / "RP1_ICSMD2026_中文初稿_v2.docx"
OUTPUT = PAPER_DIR / "RP1_ICSMD2026_中文正式稿_v4.docx"
WORK = PAPER_DIR / "_work_v4"
PIPELINE_PNG = WORK / "fig1_pipeline_v4.png"
THRESHOLD_PNG = WORK / "fig2_threshold_v4.png"

SKILL_DIR = Path(
    r"C:\Users\25021\.codex\plugins\cache\openai-primary-runtime\documents\26.805.11740\skills\documents"
)
sys.path.insert(0, str(SKILL_DIR / "scripts"))
from table_geometry import apply_table_geometry, column_widths_from_weights  # noqa: E402


def set_font(run, *, east_asia="SimSun", ascii_font="Times New Roman", size=None, bold=None, italic=None):
    run.font.name = ascii_font
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), east_asia)
    run._element.rPr.rFonts.set(qn("w:ascii"), ascii_font)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), ascii_font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def clean_inline(text: str) -> str:
    text = text.replace("`", "")
    text = text.replace("\\(", "").replace("\\)", "")
    text = text.replace("*", "")
    return text.strip()


def add_rich_text(paragraph, text: str, size=None):
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if not part:
            continue
        is_bold = part.startswith("**") and part.endswith("**")
        value = part[2:-2] if is_bold else part
        value = clean_inline(value)
        run = paragraph.add_run(value)
        set_font(run, size=size, bold=True if is_bold else None)


def remove_body_content(doc: Document):
    body = doc._element.body
    sect_pr = body.sectPr
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)


def set_columns(section, count: int, space_dxa: int):
    sect_pr = section._sectPr
    cols = sect_pr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sect_pr.append(cols)
    cols.set(qn("w:num"), str(count))
    cols.set(qn("w:space"), str(space_dxa))
    for child in list(cols):
        cols.remove(child)


def configure_sections(doc: Document):
    first = doc.sections[0]
    first.page_width = Inches(8.268)
    first.page_height = Inches(11.693)
    first.left_margin = Inches(0.62)
    first.right_margin = Inches(0.62)
    first.top_margin = Inches(0.375)
    first.bottom_margin = Inches(1.0)
    set_columns(first, 1, 720)
    first.different_first_page_header_footer = True

    footer = first.first_page_footer
    p = footer.paragraphs[0]
    p.text = "XXX-X-XXXX-XXXX-X/XX/$XX.00 ©20XX IEEE"
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in p.runs:
        set_font(run, size=8)

    second = doc.add_section(WD_SECTION.CONTINUOUS)
    second.page_width = Inches(8.268)
    second.page_height = Inches(11.693)
    second.left_margin = Inches(0.63)
    second.right_margin = Inches(0.63)
    second.top_margin = Inches(0.75)
    second.bottom_margin = Inches(1.0)
    set_columns(second, 2, 360)
    second.different_first_page_header_footer = False
    second.header.is_linked_to_previous = True
    second.footer.is_linked_to_previous = True
    return second


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def shade_cell(cell, fill="D9E2F3"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_table(doc: Document, rows: list[list[str]], table_no: int):
    captions = {
        1: "表1  全候选页面定位与来源复核",
        2: "表2  发布规则故障注入结果",
        3: "表3  固定页面提示规则及token成本",
    }
    cap = doc.add_paragraph(style="Body Text")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.keep_with_next = True
    run = cap.add_run(captions.get(table_no, f"表{table_no}"))
    set_font(run, size=8, bold=True)

    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    weights = {
        5: [1.05, 0.75, 0.9, 0.9, 1.05],
        4: [1.0, 1.9, 0.75, 0.8],
        6: [0.65, 0.8, 0.8, 0.8, 1.05, 1.15],
    }[len(rows[0])]
    widths = column_widths_from_weights(weights, 4660)
    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            run = p.add_run(clean_inline(value))
            set_font(run, size=7.2, bold=(r_idx == 0))
            if r_idx == 0:
                shade_cell(cell)
    set_repeat_table_header(table.rows[0])
    apply_table_geometry(
        table,
        widths,
        table_width_dxa=sum(widths),
        indent_dxa=55,
        cell_margins_dxa={"top": 45, "bottom": 45, "start": 55, "end": 55},
    )
    after = doc.add_paragraph(style="Body Text")
    after.paragraph_format.space_after = Pt(0)
    return table


def make_figures():
    WORK.mkdir(parents=True, exist_ok=True)
    font_path = Path(r"C:\Windows\Fonts\msyh.ttc")
    zh_font = font_manager.FontProperties(fname=str(font_path))
    plt.rcParams["axes.unicode_minus"] = False

    fig, ax = plt.subplots(figsize=(3.2, 4.0), dpi=300)
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")
    labels = [
        ("冻结语料与规则", "15份构建文档 · 1889个抽取页"),
        ("页面级候选生成", "保留原文、表格位置与物理页"),
        ("证据质量门控", "结构 · 定位 · 关系支持 · 来源"),
        ("知识断言与证据记录", "语义去重，页面证据独立保存"),
        ("双层发布", "全量审计图 → 中文应用发布图"),
    ]
    colors = ["#EAF2F8", "#E8F6F3", "#FCF3CF", "#F4ECF7", "#EBF5FB"]
    ys = [0.86, 0.68, 0.50, 0.32, 0.14]
    for i, ((title, sub), y, color) in enumerate(zip(labels, ys, colors)):
        box = FancyBboxPatch(
            (0.08, y - 0.065),
            0.84,
            0.125,
            boxstyle="round,pad=0.012,rounding_size=0.02",
            facecolor=color,
            edgecolor="#2E5D7B",
            linewidth=1.0,
        )
        ax.add_patch(box)
        ax.text(0.50, y + 0.018, title, ha="center", va="center", fontsize=9.2, fontweight="bold", fontproperties=zh_font)
        ax.text(0.50, y - 0.025, sub, ha="center", va="center", fontsize=6.7, color="#333333", fontproperties=zh_font)
        if i < len(ys) - 1:
            ax.annotate("", xy=(0.50, ys[i + 1] + 0.072), xytext=(0.50, y - 0.070), arrowprops=dict(arrowstyle="-|>", color="#2E5D7B", lw=1.0))
    fig.tight_layout(pad=0.15)
    fig.savefig(PIPELINE_PNG, bbox_inches="tight", facecolor="white")
    plt.close(fig)

    thresholds = [0.80, 0.85, 0.90, 0.925, 0.95, 0.975]
    counts = [1698, 1678, 1525, 895, 867, 36]
    fig, ax = plt.subplots(figsize=(3.2, 2.15), dpi=300)
    ax.plot(thresholds, counts, marker="o", color="#2F6FB0", linewidth=1.8, markersize=4)
    ax.fill_between(thresholds, counts, color="#DCE9F7", alpha=0.75)
    ax.set_xlabel("治理分数阈值", fontsize=7.5, fontproperties=zh_font)
    ax.set_ylabel("证据合格记录数", fontsize=7.5, fontproperties=zh_font)
    ax.tick_params(axis="both", labelsize=6.5)
    ax.grid(axis="y", linestyle="--", linewidth=0.5, alpha=0.5)
    for x, y in zip(thresholds, counts):
        ax.annotate(str(y), (x, y), textcoords="offset points", xytext=(0, 5), ha="center", fontsize=5.8)
    ax.set_ylim(0, 1900)
    fig.tight_layout(pad=0.6)
    fig.savefig(THRESHOLD_PNG, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def add_figure(doc: Document, path: Path, caption: str, width=3.05):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Inches(width))
    cap = doc.add_paragraph(style="Body Text")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.first_line_indent = Inches(0)
    cap.paragraph_format.space_after = Pt(3)
    run = cap.add_run(caption)
    set_font(run, size=8)


def add_equation(doc: Document, latex_lines: list[str]):
    raw = " ".join(line.strip() for line in latex_lines if line.strip() and "\\tag" not in line)
    equation_no = ""
    tag = next((line for line in latex_lines if "\\tag" in line), "")
    m = re.search(r"\\tag\{(\d+)\}", tag)
    if m:
        equation_no = f"({m.group(1)})"
    display = {
        "1": "cᵢ = (hᵢ, rᵢ, tᵢ)",
        "2": "aᵢ = ⟨cᵢ, eᵢ, pᵢ, gᵢ⟩",
        "3": "H(aᵢ) = Istr·Idir·Iloc·Isup·Iprov·Ibuild·Inoninf·Iq≥0.8",
        "4": "s_f(c)=max[aᵢ∈A_f(c)]qᵢ；S_fam^(B)(c)=(1/B)Σs_(j)(c)",
    }
    replacements = {
        r"\langle": "⟨",
        r"\rangle": "⟩",
        r"\in": "∈",
        r"\max": "max",
        r"\min": "min",
        r"\sum": "Σ",
        r"\geq": "≥",
        r"\qquad": "    ",
        r"\frac": "frac",
        r"\mathcal": "",
        "\\": "",
        "{": "",
        "}": "",
    }
    for old, new in replacements.items():
        raw = raw.replace(old, new)
    raw = re.sub(r"\s+", " ", raw).strip().rstrip(",.")
    if m:
        raw = display.get(m.group(1), raw)
    p = doc.add_paragraph(style="equation" if "equation" in doc.styles else "Body Text")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(raw + ("    " + equation_no if equation_no else ""))
    set_font(run, east_asia="SimSun", ascii_font="Cambria Math", size=9)
    return raw


def add_body_paragraph(doc: Document, text: str):
    p = doc.add_paragraph(style="Body Text")
    add_rich_text(p, text, size=10)
    return p


def parse_table(lines: list[str]) -> list[list[str]]:
    rows = []
    for idx, line in enumerate(lines):
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if idx == 1 and all(re.fullmatch(r":?-+:?", c) for c in cells):
            continue
        rows.append(cells)
    return rows


def build():
    make_figures()
    shutil.copy2(TEMPLATE, OUTPUT)
    doc = Document(OUTPUT)
    remove_body_content(doc)

    text = SOURCE_MD.read_text(encoding="utf-8")
    lines = text.splitlines()
    title = lines[0].removeprefix("# ").strip()
    english_title = clean_inline(lines[2])
    author = lines[4].strip()
    affiliation = lines[6].strip()

    title_p = doc.add_paragraph(style="paper title")
    title_p.text = ""
    r = title_p.add_run(title)
    set_font(r, east_asia="Microsoft YaHei", size=16, bold=True)
    en_p = doc.add_paragraph()
    en_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    en_p.paragraph_format.space_after = Pt(3)
    r = en_p.add_run(english_title)
    set_font(r, size=9, italic=True)
    author_p = doc.add_paragraph(style="Author" if "Author" in doc.styles else "Normal")
    author_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = author_p.add_run(author)
    set_font(r, size=9)
    aff_p = doc.add_paragraph(style="Affiliation" if "Affiliation" in doc.styles else "Normal")
    aff_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = aff_p.add_run(affiliation)
    set_font(r, size=8)

    abstract_idx = lines.index("## 摘要")
    kw_idx = next(i for i, line in enumerate(lines) if line.startswith("**关键词"))
    abstract_text = "".join(line.strip() for line in lines[abstract_idx + 1 : kw_idx] if line.strip())
    abs_p = doc.add_paragraph(style="Abstract")
    lead = abs_p.add_run("摘要—")
    set_font(lead, size=9, bold=True, italic=True)
    add_rich_text(abs_p, abstract_text, size=9)
    kw_p = doc.add_paragraph(style="Keywords")
    add_rich_text(kw_p, lines[kw_idx], size=9)

    configure_sections(doc)

    body_lines = lines[kw_idx + 1 :]
    i = 0
    table_no = 0
    in_refs = False
    while i < len(body_lines):
        line = body_lines[i].strip()
        if not line or line.startswith(">"):
            i += 1
            continue
        if line.startswith("## "):
            heading = line[3:].strip()
            if heading == "参考文献":
                in_refs = True
                p = doc.add_paragraph(style="Heading 5" if "Heading 5" in doc.styles else "Heading 1")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run("参考文献")
                set_font(r, size=9)
            elif heading == "致谢与声明":
                p = doc.add_paragraph(style="Heading 5" if "Heading 5" in doc.styles else "Heading 1")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run("致谢与声明")
                set_font(r, size=9)
            else:
                p = doc.add_paragraph(style="Heading 1")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                heading = re.sub(r"^[IVX]+\.\s*", "", heading)
                r = p.add_run(heading)
                set_font(r, size=10)
            i += 1
            continue
        if line.startswith("### "):
            p = doc.add_paragraph(style="Heading 2")
            subheading = re.sub(r"^[A-Z]\.\s*", "", line[4:].strip())
            r = p.add_run(subheading)
            set_font(r, size=10, italic=True)
            i += 1
            continue
        if line == r"\[":
            eq_lines = []
            i += 1
            while i < len(body_lines) and body_lines[i].strip() != r"\]":
                eq_lines.append(body_lines[i])
                i += 1
            raw = add_equation(doc, eq_lines)
            if raw.startswith("H("):
                add_figure(doc, PIPELINE_PNG, "图1  自动证据质量门控与双层发布流程", width=3.0)
            i += 1
            continue
        if line.startswith("|"):
            table_lines = []
            while i < len(body_lines) and body_lines[i].strip().startswith("|"):
                table_lines.append(body_lines[i].strip())
                i += 1
            table_no += 1
            add_table(doc, parse_table(table_lines), table_no)
            continue
        if re.match(r"^\*\*表\d+", line):
            i += 1
            continue
        if re.match(r"^- RQ\d", line):
            p = doc.add_paragraph(style="List Bullet" if "List Bullet" in doc.styles else "Body Text")
            p.paragraph_format.space_after = Pt(0.8)
            r = p.add_run(line[2:])
            set_font(r, size=10)
            i += 1
            continue
        if re.match(r"^\d+\. \*\*", line):
            text_no_num = re.sub(r"^\d+\.\s*", "", line)
            p = doc.add_paragraph(style="List Number" if "List Number" in doc.styles else "Body Text")
            p.paragraph_format.space_after = Pt(1)
            add_rich_text(p, text_no_num, size=10)
            i += 1
            continue
        if in_refs and re.match(r"^\[\d+\]", line):
            p = doc.add_paragraph(style="references")
            reference_text = re.sub(r"^\[\d+\]\s*", "", line)
            r = p.add_run(clean_inline(reference_text))
            set_font(r, size=8)
            i += 1
            continue

        p = add_body_paragraph(doc, line)
        if line.startswith("分数阈值由0.80"):
            add_figure(doc, THRESHOLD_PNG, "图2  治理分数阈值对证据合格记录数的影响", width=3.0)
        i += 1

    # A final continuous section break balances the last two-column page.
    balance = doc.add_section(WD_SECTION.CONTINUOUS)
    balance.page_width = Inches(8.268)
    balance.page_height = Inches(11.693)
    balance.left_margin = Inches(0.63)
    balance.right_margin = Inches(0.63)
    balance.top_margin = Inches(0.75)
    balance.bottom_margin = Inches(1.0)
    set_columns(balance, 1, 360)
    balance.header.is_linked_to_previous = True
    balance.footer.is_linked_to_previous = True

    # Keep headings and captions with following content.
    for p in doc.paragraphs:
        if p.style.name in {"Heading 1", "Heading 2", "Heading 5"}:
            p.paragraph_format.keep_with_next = True
        for run in p.runs:
            color = run.font.color.rgb
            if color is not None and color != RGBColor(0, 0, 0):
                run.font.color.rgb = RGBColor(0, 0, 0)

    doc.core_properties.title = title
    doc.core_properties.subject = "ICSMD 2026 Chinese formal manuscript"
    doc.core_properties.keywords = "marine pump; knowledge graph; evidence quality gating; provenance"
    doc.core_properties.comments = "Chinese formal draft; final submission must be translated into English."
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
