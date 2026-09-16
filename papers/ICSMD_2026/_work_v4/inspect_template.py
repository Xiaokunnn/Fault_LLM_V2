from collections import Counter
from pathlib import Path
from zipfile import ZipFile
import hashlib
from lxml import etree

from docx import Document
from docx.oxml.ns import qn


def emu_in(value):
    return None if value is None else round(value.inches, 3)


def pt(value):
    return None if value is None else round(value.pt, 2)


path = Path(__file__).resolve().parents[1] / "RP1_ICSMD2026_中文初稿_v2.docx"
doc = Document(path)
print("FILE", path)
print("SHA256", hashlib.sha256(path.read_bytes()).hexdigest().upper())
print("SECTIONS", len(doc.sections))
for i, sec in enumerate(doc.sections, 1):
    cols = sec._sectPr.find(qn("w:cols"))
    print(
        "SECTION",
        i,
        "size",
        emu_in(sec.page_width),
        emu_in(sec.page_height),
        "margins",
        emu_in(sec.left_margin),
        emu_in(sec.right_margin),
        emu_in(sec.top_margin),
        emu_in(sec.bottom_margin),
        "columns",
        None if cols is None else etree.tostring(cols, encoding="unicode"),
    )

counts = Counter(p.style.name for p in doc.paragraphs)
print("STYLE_COUNTS", counts)
roles = [
    "paper title",
    "author",
    "Affiliation",
    "Abstract",
    "Keywords",
    "Heading 1",
    "Heading 2",
    "Heading 3",
    "Heading 4",
    "Heading 5",
    "Normal",
    "figurecaption",
    "table caption",
    "references",
]
for name in roles:
    if name not in doc.styles:
        continue
    st = doc.styles[name]
    pf = st.paragraph_format
    font = st.font
    rpr = st.element.rPr
    fonts = None if rpr is None else rpr.rFonts
    print(
        "STYLE",
        name,
        "font",
        font.name,
        "eastAsia",
        None if fonts is None else fonts.get(qn("w:eastAsia")),
        "size",
        pt(font.size),
        "bold",
        font.bold,
        "italic",
        font.italic,
        "align",
        pf.alignment,
        "before",
        pt(pf.space_before),
        "after",
        pt(pf.space_after),
        "line",
        pf.line_spacing,
        "first",
        emu_in(pf.first_line_indent),
        "left",
        emu_in(pf.left_indent),
    )

print("PARAGRAPHS", len(doc.paragraphs), "TABLES", len(doc.tables), "IMAGES", len(doc.inline_shapes))
with ZipFile(path) as zf:
    for info in sorted(zf.infolist(), key=lambda x: x.filename):
        data = zf.read(info.filename)
        print("PART", info.filename, info.file_size, hashlib.sha256(data).hexdigest().upper())
